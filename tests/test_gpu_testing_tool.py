import unittest
import tempfile
import threading
import time
from pathlib import Path

import numpy as np

from gpu_testing_tool import (
    InferenceOutcome,
    InferenceResult,
    SessionWorker,
    prepare_image_for_model,
    run_direct_concurrent_groups,
    run_inference_batch,
    run_worker_concurrent_groups,
    run_two_sender_jobs,
    resolve_numpy_dtype,
    ScenarioResult,
    get_environment_info,
    _is_oom_error,
    write_results_to_excel,
)


class PrepareImageTests(unittest.TestCase):
    def test_preserves_dynamic_height_and_width(self):
        arr = np.arange(6, dtype=np.uint16).reshape(2, 3)
        out = prepare_image_for_model(arr, [1, None, None, 1], ["batch", "height", "width", "channel"])
        self.assertEqual(out.shape, (1, 2, 3, 1))
        np.testing.assert_array_equal(out[0, :, :, 0], arr)

    def test_rejects_fixed_spatial_dimension_mismatch(self):
        arr = np.arange(24, dtype=np.uint16).reshape(4, 6)
        with self.assertRaisesRegex(ValueError, "never cropped or padded"):
            prepare_image_for_model(arr, [1, 3, 5, 1], ["batch", "height", "width", "channel"])

    def test_rejects_models_without_named_spatial_axes(self):
        with self.assertRaisesRegex(ValueError, "height and one width"):
            prepare_image_for_model(np.zeros((2, 3)), [1, None, None, 1], ["batch", None, None, "channel"])

    def test_accepts_static_single_channel_hwc_without_dimension_names(self):
        arr = np.arange(6, dtype=np.int32).reshape(2, 3)
        out = prepare_image_for_model(arr, [2, 3, 1], [None, None, None])
        self.assertEqual(out.shape, (2, 3, 1))
        np.testing.assert_array_equal(out[:, :, 0], arr)


class DTypeTests(unittest.TestCase):
    def test_resolve_uint16(self):
        self.assertEqual(resolve_numpy_dtype("tensor(uint16)"), np.dtype(np.uint16))


class FailingSession:
    def run(self, *_args, **_kwargs):
        raise RuntimeError("CUDA out of memory")


class ConcurrentSession:
    def __init__(self):
        self.active = 0
        self.max_active = 0
        self.lock = threading.Lock()

    def run(self, *_args, **_kwargs):
        with self.lock:
            self.active += 1
            self.max_active = max(self.max_active, self.active)
        time.sleep(0.03)
        with self.lock:
            self.active -= 1
        return []


class RecordingWorker:
    def __init__(self):
        self.submit_times = []

    def submit(self, image_name, _arr, result_queue):
        self.submit_times.append((image_name, time.perf_counter()))
        result_queue.put(InferenceOutcome(result=InferenceResult(image_name, 0.0)))


class TimedRecordingWorker(RecordingWorker):
    def submit(self, image_name, _arr, result_queue, on_start=None):
        if on_start is not None:
            on_start()
        super().submit(image_name, _arr, result_queue)


class ExecutionTests(unittest.TestCase):
    def test_bad_allocation_is_classified_as_oom(self):
        self.assertTrue(_is_oom_error(RuntimeError("Concat failed: bad allocation")))

    def test_worker_exception_is_returned_to_caller(self):
        worker = SessionWorker(0, FailingSession(), "input")
        worker.start()
        try:
            with self.assertRaisesRegex(RuntimeError, "CUDA out of memory"):
                run_inference_batch([worker], [(0, "image", np.zeros((1,), dtype=np.float32))])
        finally:
            worker.stop()
            worker.join(timeout=1)
        self.assertFalse(worker.is_alive())

    def test_direct_concurrent_groups_call_one_session_concurrently(self):
        session = ConcurrentSession()
        images = [(str(i), np.zeros((1,), dtype=np.float32), (1,)) for i in range(3)]
        results = run_direct_concurrent_groups(session, "input", images, concurrency=3)
        self.assertEqual(len(results), 1)
        self.assertGreaterEqual(session.max_active, 2)

    def test_worker_groups_use_first_worker_start_as_request_start(self):
        workers = [TimedRecordingWorker(), TimedRecordingWorker()]
        images = [(str(i), np.zeros((1,), dtype=np.float32), (1,)) for i in range(3)]
        results = run_worker_concurrent_groups(workers, images, task_concurrency=2)
        self.assertEqual([result.image_name for result in results], ["request_group_1", "request_group_2"])
        self.assertEqual(len(workers[0].submit_times), 2)
        self.assertEqual(len(workers[1].submit_times), 1)

    def test_two_senders_keep_fixed_submission_cadence(self):
        workers = [RecordingWorker(), RecordingWorker()]
        jobs = [
            [(0, "a", np.zeros(1)), (0, "c", np.zeros(1))],
            [(1, "b", np.zeros(1)), (1, "d", np.zeros(1))],
        ]
        run_two_sender_jobs(workers, jobs, interval_ms=30)
        for worker in workers:
            self.assertGreaterEqual(worker.submit_times[1][1] - worker.submit_times[0][1], 0.02)

    def test_scenario4_uses_request_group_headers(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        result = ScenarioResult("scenario4a_test", "details", [], 0.0, 0.0)
        env_info = {
            "cuda_version": "x",
            "tensorrt_version": "x",
            "model_size_bytes": 1,
            "image_size_bytes": 1,
            "original_image_shape": "(100, 200)",
            "model_input_shape": "(1, 1, 100, 200)",
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "results.xlsx"
            write_results_to_excel(str(output), env_info, [result])
            sheet = openpyxl.load_workbook(output).active
            self.assertEqual(sheet.cell(row=12, column=1).value, "request_group")
            self.assertEqual(sheet.cell(row=7, column=1).value, "original_image_shape")
            self.assertEqual(sheet.cell(row=7, column=2).value, "(100, 200)")


if __name__ == "__main__":
    unittest.main()
